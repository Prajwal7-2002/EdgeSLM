"""Pull plain text out of the file types a company drive is full of.

Everything here works from bytes plus a file name, so the API can classify an
upload without writing it to disk. Scanned PDFs (images, no text layer) come
back empty and are rejected rather than classified as an empty document.
"""

import email
import io
from email import policy
from html.parser import HTMLParser
from pathlib import Path


TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".log", ".json", ".yaml", ".yml", ".ini", ".cfg",
    ".conf", ".env", ".xml", ".py", ".js", ".ts", ".sh", ".ps1", ".sql",
    ".pem", ".key", ".crt", ".pub", ".cer",
}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".pdf", ".docx", ".eml", ".html", ".htm"}


class UnsupportedFile(ValueError):
    pass


def _pdf(data):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(data):
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # Credentials and salary figures often live in tables, not paragraphs.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _eml(data):
    message = email.message_from_bytes(data, policy=policy.default)
    parts = [f"Subject: {message['subject'] or ''}"]
    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        parts.append(body.get_content())
    return "\n\n".join(parts)


class _TextFromHtml(HTMLParser):
    SKIP = {"script", "style", "head"}
    BLOCKS = {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "table", "section"}

    def __init__(self):
        super().__init__()
        self.parts = []
        self.skipping = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skipping:
            self.skipping -= 1

    def handle_data(self, data):
        if not self.skipping:
            self.parts.append(data)


def _html(data):
    parser = _TextFromHtml()
    parser.feed(data.decode("utf-8", errors="replace"))
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    return "\n".join(line for line in lines if line)


def extract_text(data, filename):
    """Return the text of a file given its bytes and name."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _pdf(data)
    elif suffix == ".docx":
        text = _docx(data)
    elif suffix == ".eml":
        text = _eml(data)
    elif suffix in (".html", ".htm"):
        text = _html(data)
    elif suffix in TEXT_SUFFIXES:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFile(
            f"unsupported file type {suffix or '(none)'}; "
            f"supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )

    text = text.strip()
    if not text:
        raise UnsupportedFile(f"{filename} has no extractable text (a scanned PDF needs OCR first)")
    return text


def extract_file(path):
    path = Path(path)
    return extract_text(path.read_bytes(), path.name)
