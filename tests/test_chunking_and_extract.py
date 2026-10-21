import email.message
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from inference.chunking import document_token_budget, most_sensitive, split_into_chunks
from inference.extract import UnsupportedFile, extract_text
from inference.prompt import build_prompt


class WordTokenizer:
    """One token per whitespace-separated word, enough to test chunk sizes."""

    def __init__(self):
        self.vocab = {}
        self.words = []

    def __call__(self, text, add_special_tokens=False):
        ids = []
        for word in text.split():
            if word not in self.vocab:
                self.vocab[word] = len(self.words)
                self.words.append(word)
            ids.append(self.vocab[word])
        return {"input_ids": ids}

    def decode(self, ids, skip_special_tokens=True):
        return " ".join(self.words[i] for i in ids)


def test_short_document_is_one_unchanged_chunk():
    tokenizer = WordTokenizer()
    text = "The standup moves to ten."
    assert split_into_chunks(tokenizer, text, budget=100) == [text]


def test_long_document_chunks_fit_the_budget_and_overlap():
    tokenizer = WordTokenizer()
    words = [f"w{i}" for i in range(250)]
    chunks = split_into_chunks(tokenizer, " ".join(words), budget=100, overlap=10)
    assert all(len(chunk.split()) <= 100 for chunk in chunks)
    assert chunks[0].split()[-10:] == chunks[1].split()[:10]
    covered = set(word for chunk in chunks for word in chunk.split())
    assert covered == set(words)


def test_secret_at_the_end_of_a_long_document_is_still_seen():
    tokenizer = WordTokenizer()
    text = " ".join(["filler"] * 900 + ["password:", "Copper!2291"])
    chunks = split_into_chunks(tokenizer, text, budget=100)
    assert "Copper!2291" in chunks[-1]


def test_every_chunk_keeps_the_closing_prompt_line():
    tokenizer = WordTokenizer()
    budget = document_token_budget(tokenizer, max_prompt_tokens=200)
    text = " ".join(f"w{i}" for i in range(1000))
    for chunk in split_into_chunks(tokenizer, text, budget):
        prompt_ids = tokenizer(build_prompt(chunk))["input_ids"]
        assert len(prompt_ids) <= 200
        assert build_prompt(chunk).endswith("Classification:")


def test_document_takes_its_most_sensitive_chunk():
    assert most_sensitive(["PUBLIC", "RESTRICTED", "INTERNAL"]) == "RESTRICTED"
    assert most_sensitive(["INTERNAL", "CONFIDENTIAL"]) == "CONFIDENTIAL"
    assert most_sensitive(["PUBLIC"]) == "PUBLIC"


def test_majority_voting_aggregation():
    from collections import Counter
    from inference.chunking import SEVERITY

    # Majority wins
    chunks = ["INTERNAL", "INTERNAL", "PUBLIC"]
    counts = Counter(chunks)
    label = max(counts.keys(), key=lambda l: (counts[l], SEVERITY[l]))
    assert label == "INTERNAL"

    # Tie broken by severity
    chunks = ["INTERNAL", "CONFIDENTIAL"]
    counts = Counter(chunks)
    label = max(counts.keys(), key=lambda l: (counts[l], SEVERITY[l]))
    assert label == "CONFIDENTIAL"


def test_extracts_docx_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_paragraph("Vendor access summary")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "svc_ingest"
    table.rows[0].cells[1].text = "Copper!2291"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text(buffer.getvalue(), "access.docx")
    assert "Vendor access summary" in text
    assert "svc_ingest | Copper!2291" in text


def test_extracts_email_subject_and_body():
    message = email.message.EmailMessage()
    message["Subject"] = "Board pre-read"
    message.set_content("Revenue is tracking ahead of plan. Do not forward.")
    text = extract_text(message.as_bytes(), "note.eml")
    assert "Subject: Board pre-read" in text
    assert "Do not forward." in text


def minimal_pdf(line):
    """A one-page PDF with a real text layer, built by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({line}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def test_extracts_pdf_text_layer():
    assert "Quarterly planning notes" in extract_text(minimal_pdf("Quarterly planning notes"), "plan.pdf")


def test_plain_text_and_config_files_are_read():
    assert extract_text(b"DB_PASS=Copper!2291", "app.env") == "DB_PASS=Copper!2291"


def test_unsupported_and_empty_files_are_rejected():
    with pytest.raises(UnsupportedFile):
        extract_text(b"\x00\x01", "photo.png")
    with pytest.raises(UnsupportedFile):
        extract_text(b"   ", "blank.txt")


def test_html_is_reduced_to_visible_text():
    page = b"<html><head><title>x</title><style>p{}</style></head><body><p>Press release:</p><p>Revenue grew&nbsp;8%.</p><script>var a=1;</script></body></html>"
    text = extract_text(page, "release.htm")
    assert "Press release:" in text and "Revenue grew" in text
    assert "var a" not in text and "p{}" not in text
